import docx
import io
import sys
from fastapi.testclient import TestClient
from app import app

sys.stdout.reconfigure(encoding='utf-8')

client = TestClient(app)

def run_tests():
    # 1. Test homepage
    r = client.get('/')
    assert r.status_code == 200, f"Status {r.status_code}"
    assert "humanize studio" in r.text.lower()
    print("✓ GET / homepage OK")

    # 2. Test text humanize
    sample_ai = "Furthermore, it is crucial to delve into the intricate tapestry of modern artificial intelligence. In conclusion, it serves as a testament to human innovation."
    r2 = client.post('/api/humanize-text', json={'text': sample_ai, 'mode': 'ultra_stealth', 'intensity': 75})
    assert r2.status_code == 200, f"Status {r2.status_code}"
    data = r2.json()
    print("✓ POST /api/humanize-text OK")
    print("  Original Human %:", data['metrics_before']['human_score'])
    print("  Humanized Human %:", data['metrics_after']['human_score'])
    print("  Humanized Text:", data['humanized_text'])

    # 3. Test single docx upload
    doc = docx.Document()
    doc.add_heading('AI In Healthcare', 1)
    doc.add_paragraph('Moreover, machine learning algorithms play a pivotal role in diagnostics. It is crucial to note that doctors utilize these models.')
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    r3 = client.post(
        '/api/upload-document',
        files={'file': ('healthcare.docx', buf.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},
        data={'mode': 'academic', 'intensity': '75'}
    )
    assert r3.status_code == 200, f"Doc status {r3.status_code}"
    doc_data = r3.json()
    token = doc_data['token']
    print("✓ POST /api/upload-document OK, Token:", token)

    # 4. Test download
    r_down = client.get(f"/api/download/{token}")
    assert r_down.status_code == 200
    assert len(r_down.content) > 100
    print("✓ GET /api/download/{token} OK, Downloaded bytes:", len(r_down.content))

    # 5. Test batch upload
    buf2 = io.BytesIO()
    doc.save(buf2)
    r_batch = client.post(
        '/api/upload-batch',
        files=[
            ('files', ('doc1.docx', buf.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')),
            ('files', ('doc2.docx', buf2.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'))
        ],
        data={'mode': 'ultra_stealth', 'intensity': '75'}
    )
    assert r_batch.status_code == 200
    batch_data = r_batch.json()
    print("✓ POST /api/upload-batch OK, ZIP Token:", batch_data['token'], "Files processed:", batch_data['total_files'])

    # 6. Test zip download
    zip_token = batch_data['token']
    r_zip = client.get(f"/api/download/{zip_token}")
    assert r_zip.status_code == 200
    assert len(r_zip.content) > 200
    print("✓ Batch ZIP Download OK, Size:", len(r_zip.content), "bytes")
    print("=========================================================")
    print("🎉 ALL INTEGRATION TESTS PASSED 100% PERFECTLY!")
    print("=========================================================")

if __name__ == '__main__':
    run_tests()
